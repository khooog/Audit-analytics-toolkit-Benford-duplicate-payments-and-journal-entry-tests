"""
Independent check of the audit analytics workbook.
1. Recalculates the workbook with LibreOffice (headless) so every formula is evaluated.
2. Regenerates the sample data and runs the pandas implementation of every test.
3. Compares every flag, score and Benford statistic row by row with the workbook.
4. Scans for formula errors and reads the Checks sheet.
Run: python verify_workbook.py path/to/audit_analytics_workbook.xlsx
Needs: LibreOffice (soffice) on the PATH, pandas, openpyxl.
"""
import os
import subprocess
import sys
import tempfile

from openpyxl import load_workbook

import audit_analytics as aa

ERRS = ("#REF!", "#NAME?", "#VALUE!", "#DIV/0!", "#N/A", "#NUM!", "#NULL!")


def sheet_rows(ws, first, cols):
    out = {}
    for row in ws.iter_rows(min_row=first, values_only=True):
        if row[0] is None:
            continue
        out[row[0]] = row
    return out


def main(path):
    tmp = tempfile.mkdtemp()
    subprocess.run(["soffice", "--headless", "--convert-to", "xlsx", "--outdir", tmp, path], check=True, capture_output=True, timeout=500)
    wb = load_workbook(os.path.join(tmp, os.path.basename(path)), data_only=True)
    errors = [(ws.title, c.coordinate, c.value) for ws in wb.worksheets for row in ws.iter_rows() for c in row
              if isinstance(c.value, str) and c.value in ERRS]
    print("Formula errors found:", errors if errors else "none")
    ok = not errors

    data = aa.gen_data()
    p, j, ben, vb = aa.pandas_tests(data)

    # Payments
    PCOLS = ["id", "date", "vendor", "inv", "amt", "appr", "fd", "norm", "dup", "near", "rnd", "below", "split", "auth", "emp", "benf", "score", "pri"]
    ws = wb["Payments"]
    rows = sheet_rows(ws, 5, PCOLS)
    bad = 0
    for r in p.itertuples():
        w = dict(zip(PCOLS, rows[r.id]))
        for k in ["fd", "dup", "near", "rnd", "below", "split", "auth", "emp", "benf", "score"]:
            if abs(float(w[k]) - float(getattr(r, k))) > 1e-9:
                bad += 1
                print("Payment mismatch", r.id, k, w[k], getattr(r, k))
        if (w["pri"] or "") != r.pri:
            bad += 1
    print("Payments compared:", len(p), "mismatches:", bad)
    ok &= bad == 0

    # Journal entries
    JC = ["no", "n", "debit", "credit", "diff", "post", "doc", "user", "source", "desc", "lag", "rev", "susp_n",
          "unb", "wkd", "late", "pend", "unauth", "mrev", "rnd", "nodesc", "below", "susp", "gap", "score", "pri"]
    ws = wb["JE_Entries"]
    rows = sheet_rows(ws, 5, JC)
    bad = 0
    for r in j.itertuples():
        w = dict(zip(JC, rows[int(r.no)]))
        for k, pk in [("unb", "unb"), ("wkd", "wkd"), ("late", "late"), ("pend", "pend"), ("unauth", "unauth"), ("mrev", "mrev"),
                      ("rnd", "rnd"), ("nodesc", "nodesc"), ("below", "below"), ("susp", "susp_f"), ("gap", "gap"), ("score", "score")]:
            if abs(float(w[k]) - float(getattr(r, pk))) > 1e-9:
                bad += 1
                print("JE mismatch", r.no, k, w[k], getattr(r, pk))
        if abs(float(w["debit"]) - r.debit) > 0.005 or abs(float(w["credit"]) - r.credit) > 0.005:
            bad += 1
        if (w["pri"] or "") != r.pri:
            bad += 1
    print("Journal entries compared:", len(j), "mismatches:", bad)
    ok &= bad == 0

    # Benford
    ws = wb["Benford"]
    n, mad, chi = ws["B15"].value, ws["B16"].value, ws["B18"].value
    zs = [ws.cell(row=4 + d, column=6).value for d in range(1, 10)]
    zdiff = max(abs(a - b) for a, b in zip(zs, ben["z"]))
    good = n == ben["n"] and abs(mad - ben["mad"]) < 1e-9 and abs(chi - ben["chi"]) < 1e-6 and zdiff < 1e-6
    print(f"Benford n={n} MAD={mad:.6f} (python {ben['mad']:.6f}) chi={chi:.4f} (python {ben['chi']:.4f}) max z diff={zdiff:.2e}", "OK" if good else "FAIL")
    ok &= good
    vws = wb["Vendor_Benford"]
    vbad = 0
    for row in vws.iter_rows(min_row=6, values_only=True):
        if row[0] in vb:
            v = vb[row[0]]
            if int(row[1]) != v["n"] or int(row[12]) != v["flag"] or int(row[13]) != v["emp"]:
                vbad += 1
            if v["mad"] is not None and abs(row[11] - v["mad"]) > 1e-9:
                vbad += 1
    print("Vendor Benford mismatches:", vbad)
    ok &= vbad == 0

    # Checks sheet and exceptions
    ws = wb["Checks"]
    for row in ws.iter_rows(min_row=5, values_only=True):
        if row[0] and row[3] in ("OK", "FAIL", "PASS"):
            print(f"  {row[0][:78]:78s} {row[1]!s:>8} {row[3]}")
    overall = [c.value for row in ws.iter_rows() for c in row if c.value in ("PASS", "FAIL")]
    print("Checks overall:", overall)
    ok &= overall == ["PASS"]
    ws = wb["Exceptions"]
    print("Test key | flagged | seeded | detected | recall | flagged not in key")
    for row in ws.iter_rows(min_row=5, max_row=23, values_only=True):
        print(f"  {row[0]:9s} {row[3]:>4} {row[5]:>4} {row[6]:>4} {row[7]!s:>6} {row[8]:>4}")
    print("ALL VERIFICATIONS PASSED" if ok else "VERIFICATION FAILED")
    return 0 if ok else 1


if __name__ == "__main__":
    sys.exit(main(sys.argv[1]))
