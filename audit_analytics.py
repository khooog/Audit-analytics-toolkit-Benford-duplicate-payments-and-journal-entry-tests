"""
Audit analytics toolkit: Benford's Law, duplicate payments, round sums,
threshold avoidance and journal entry tests, with risk scoring.

ALL DATA IN THIS FILE IS SAMPLE / DEMONSTRATION DATA generated with a fixed seed.
No real company, vendor, employee or transaction is represented. Anomalies are
seeded on purpose so the tests can be validated against an answer key.

Run:  python audit_analytics.py [output_folder]
Needs: Python 3.9+, pandas, openpyxl

Output is a list of items for targeted follow up. It is not an audit sample
and results must not be projected to the population (see ISA 530 on selecting
specific items versus audit sampling).
"""
import datetime as dt
import json
import math
import os
import random
import sys
from collections import defaultdict

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

SEED = 20260
FY_START, FY_END = dt.date(2025, 1, 1), dt.date(2025, 12, 31)
SAMPLE_LABEL = "Sample / demonstration data only. No real company, vendor, employee or transaction."

# ----------------------------------------------------------------------------
# Parameters (single source for the engine and the workbook Parameters sheet)
# ----------------------------------------------------------------------------
PARAMS = [
    ("fy_start", "Fiscal year start", FY_START, "Population period"),
    ("fy_end", "Fiscal year end", FY_END, "Population period"),
    ("pay_limit", "Payment approval limit (USD)", 5000, "Payments at or above need the senior approver"),
    ("senior", "Senior approver code", "FIN_DIR", "Approver required at or above the limit"),
    ("pay_band", "Below limit band", 0.05, "Flag payments within this share below the limit"),
    ("round_mod", "Round sum modulus (USD)", 1000, "Amount is a multiple of this"),
    ("round_min", "Round sum minimum (USD)", 1000, "Ignore smaller round amounts"),
    ("dup_days", "Near duplicate window (days)", 30, "Same vendor and amount within this many days"),
    ("ben_minn", "Benford minimum payments per vendor", 30, "Vendor test is unreliable below this count"),
    ("ben_mad", "Benford MAD threshold", 0.015, "First digit MAD above this is nonconformity (Nigrini)"),
    ("ben_z", "Benford z critical value", 1.96, "5 percent two sided"),
    ("ben_chi", "Chi square critical value (8 df, 5 percent)", 15.507, "Fixed statistical constant"),
    ("je_lag", "JE posting lag limit (days)", 15, "Posting date less document date"),
    ("je_end", "Period end window (days)", 3, "Last days of the fiscal year"),
    ("je_limit", "JE approval limit (USD)", 50000, "Manual entries need extra approval at or above"),
    ("je_band", "JE below limit band", 0.05, "Flag manual entries within this share below the limit"),
    ("je_round_mod", "JE round modulus (USD)", 1000, "Entry total is a multiple of this"),
    ("je_round_min", "JE round minimum (USD)", 10000, "Ignore smaller round entries"),
    ("tol", "Balance tolerance (USD)", 0.005, "Debit less credit beyond this is unbalanced"),
    ("cut_pay", "High risk cutoff, payments (score)", 5, "Score at or above is High"),
    ("cut_je", "High risk cutoff, journal entries (score)", 5, "Score at or above is High"),
]
PV = {k: v for k, _, v, _ in PARAMS}

PAY_W = [("dup", "Duplicate payment (same vendor, invoice, amount)", 5),
         ("near", "Near duplicate (same vendor and amount, different invoice)", 3),
         ("rnd", "Round sum", 2),
         ("below", "Just below approval limit", 2),
         ("split", "Same day split above limit", 3),
         ("auth", "Approved below required authority", 4),
         ("emp", "Vendor bank account matches an employee", 5),
         ("benf", "Vendor fails Benford test", 2)]
JE_W = [("unb", "Unbalanced entry", 5),
        ("wkd", "Posted on weekend", 2),
        ("late", "Posted late versus document date", 2),
        ("pend", "Posted in last days of year", 1),
        ("unauth", "Posted by user not authorised", 3),
        ("mrev", "Manual entry to revenue", 3),
        ("rnd", "Round total", 2),
        ("nodesc", "No description", 2),
        ("below", "Manual entry just below approval limit", 3),
        ("susp", "Manual entry to suspense", 2),
        ("gap", "Number sequence gap before entry", 1)]
for k, _, w in PAY_W:
    PV["wp_" + k] = w
for k, _, w in JE_W:
    PV["wj_" + k] = w

AUTH_USERS = ["FIN01", "FIN02", "FIN03", "FIN04", "FIN05", "SYS_AP", "SYS_AR", "SYS_PAY"]

COA = [(1000, "Cash at bank", "Asset"), (1100, "Trade receivables", "Asset"), (1200, "Inventory", "Asset"),
       (1500, "Fixed assets", "Asset"), (2000, "Trade payables", "Liability"), (2100, "Accruals", "Liability"),
       (2500, "Loans", "Liability"), (4000, "Sales revenue", "Revenue"), (4100, "Other income", "Revenue"),
       (5000, "Cost of sales", "Expense"), (6000, "Salaries", "Expense"), (6100, "Rent", "Expense"),
       (6200, "Utilities", "Expense"), (6900, "Sundry expense", "Expense"), (9999, "Suspense", "Suspense")]
CLASS = {a: c for a, _, c in COA}


# ----------------------------------------------------------------------------
# Data generation
# ----------------------------------------------------------------------------
def weekday_on_or_after(d, rnd=None, allow_weekend=False):
    while not allow_weekend and d.weekday() >= 5:
        d += dt.timedelta(days=1)
    return d


def rand_weekday(R):
    while True:
        d = FY_START + dt.timedelta(days=R.randint(0, 364))
        if d.weekday() < 5:
            return d


def rand_weekend(R):
    while True:
        d = FY_START + dt.timedelta(days=R.randint(0, 364))
        if d.weekday() >= 5:
            return d


def gen_data():
    R = random.Random(SEED)
    vendors = [f"V{i:03d}" for i in range(1, 46)]
    employees = [f"E{i:03d}" for i in range(1, 26)]
    emp_bank = {e: f"ACC{(i * 15485863) % 10 ** 8:08d}" for i, e in enumerate(employees, start=1)}
    ven_bank = {v: f"ACC{(i * 104729 + 55) % 10 ** 8:08d}" for i, v in enumerate(vendors, start=1)}
    ven_bank["V038"] = emp_bank["E014"]  # seeded: vendor bank account equals employee bank account

    counter = defaultdict(int)

    def next_inv(v):
        counter[v] += 1
        return f"INV{int(v[1:]) * 1000 + counter[v]}"

    pays = []

    def add(date, v, inv, amt, appr, tags=None, note=""):
        pays.append(dict(date=date, vendor=v, inv=inv, amt=round(amt, 2), appr=appr, tags=tags or [], note=note))

    special = {"V001", "V002", "V031"}
    organic_v = [v for v in vendors if v not in special]
    weights = [R.uniform(0.3, 3.0) for _ in organic_v]
    for _ in range(520):
        v = R.choices(organic_v, weights)[0]
        amt = 10 ** R.uniform(2, 5)
        appr = "FIN_DIR" if amt >= PV["pay_limit"] else R.choice(["AP_MGR", "AP_MGR", "FIN_DIR"])
        add(rand_weekday(R), v, next_inv(v), amt, appr)

    # recurring vendors (benign by construction)
    for m in range(1, 13):
        d = weekday_on_or_after(dt.date(2025, m, 1) + dt.timedelta(days=R.randint(0, 5)))
        add(d, "V001", f"RENT25{m:02d}", 85000.00, "FIN_DIR", [("p_rnd", "Benign by construction")], "Monthly rent")
        d = weekday_on_or_after(dt.date(2025, m, 15) + dt.timedelta(days=R.randint(0, 3)))
        add(d, "V002", f"LIC25{m:02d}", 4900.00, "AP_MGR", [("p_below", "Benign by construction")], "Monthly licence")

    # seeded duplicates built from distinct organic originals dated before December
    orig_pool = [p for p in pays[:520] if 200 <= p["amt"] <= 40000 and p["date"] < dt.date(2025, 12, 1)]
    originals = R.sample(orig_pool, 14)
    for p in originals[:6]:
        d = weekday_on_or_after(p["date"] + dt.timedelta(days=R.randint(3, 20)))
        add(d, p["vendor"], p["inv"], p["amt"], p["appr"], [("p_dup", "Seeded anomaly")], "Exact duplicate")
    variants = ["INV {n}", "inv/{n}", "INV.{n}", "inv {n}"]
    for p, f in zip(originals[6:10], variants):
        d = weekday_on_or_after(p["date"] + dt.timedelta(days=R.randint(3, 20)))
        inv = f.format(n=p["inv"][3:])
        add(d, p["vendor"], inv, p["amt"], p["appr"], [("p_dup", "Seeded anomaly")], "Duplicate, invoice number reformatted")
    for p in originals[10:14]:
        d = weekday_on_or_after(p["date"] + dt.timedelta(days=R.randint(2, 25)))
        p["tags"].append(("p_near", "Counterpart of seeded anomaly"))
        add(d, p["vendor"], next_inv(p["vendor"]), p["amt"], p["appr"], [("p_near", "Seeded anomaly")], "Same amount, new invoice number")

    for _ in range(12):  # round sums
        v = R.choice(organic_v)
        amt = 1000 * R.randint(2, 48)
        add(rand_weekday(R), v, next_inv(v), amt, "FIN_DIR" if amt >= PV["pay_limit"] else "AP_MGR",
            [("p_rnd", "Seeded anomaly")], "Round sum")

    for _ in range(45):  # threshold avoidance vendor
        add(rand_weekday(R), "V031", next_inv("V031"), R.uniform(4700, 4990), "AP_MGR",
            [("p_below", "Seeded anomaly")], "Threshold avoidance vendor")

    for _ in range(5):  # same day splits
        v = R.choice(organic_v)
        d = rand_weekday(R)
        while True:
            n = R.choice([2, 3])
            parts = [round(R.uniform(2200, 3400), 2) for _ in range(n)]
            if sum(parts) >= PV["pay_limit"]:
                break
        for a in parts:
            add(d, v, next_inv(v), a, "AP_MGR", [("p_split", "Seeded anomaly")], "Same day split")

    for _ in range(3):  # approval authority breach
        v = R.choice(organic_v)
        add(rand_weekday(R), v, next_inv(v), R.uniform(5000, 30000), "AP_MGR", [("p_auth", "Seeded anomaly")],
            "Approved below required authority")

    pays.sort(key=lambda p: (p["date"], p["vendor"], p["inv"], p["amt"]))
    for i, p in enumerate(pays, start=1):
        p["id"] = f"P{i:04d}"

    # ---------------- journal entries ----------------
    users_fin = ["FIN01", "FIN02", "FIN03", "FIN04", "FIN05"]
    entries = []

    def entry(post, doc, user, source, desc, lines, tags=None):
        entries.append(dict(post=post, doc=doc, user=user, source=source, desc=desc, lines=lines, tags=tags or []))

    def amount():
        return round(10 ** R.uniform(2.3, 5.2), 2)

    def two_lines(dr, cr, amt):
        if R.random() < 0.2:
            first = round(amt * 0.7, 2)
            return [(dr, first, 0.0), (dr, round(amt - first, 2), 0.0), (cr, 0.0, amt)]
        return [(dr, amt, 0.0), (cr, 0.0, amt)]

    templates = [
        (1100, 4000, 25, "SYS_AR", "Sales invoice batch"), (1000, 1100, 20, "SYS_AR", "Customer receipts"),
        (5000, 1200, 12, "SYS_AR", "Cost of sales posting"), (1200, 2000, 15, "SYS_AP", "Supplier invoices"),
        (2000, 1000, 20, "SYS_PAY", "Supplier payments"), (6200, 2000, 5, "SYS_AP", "Utilities invoices"),
        (1500, 2000, 3, "SYS_AP", "Fixed asset additions"), (6900, 2100, 5, "MANUAL", "Accrual for sundry costs"),
        (2100, 6900, 4, "MANUAL", "Accrual reversal"),
    ]
    tw = [t[2] for t in templates]
    for _ in range(730):
        dr, cr, _, src, desc = R.choices(templates, tw)[0]
        amt = amount()
        if src == "MANUAL":
            post = rand_weekday(R)
            entry(post, post - dt.timedelta(days=R.randint(0, 4)), R.choice(users_fin), "Manual", desc, two_lines(dr, cr, amt))
        else:
            tags = []
            if R.random() < 0.025:
                post = rand_weekend(R)
                tags.append(("j_wkd", "Benign by construction"))
            else:
                post = rand_weekday(R)
            d = desc
            if R.random() < 0.015:
                d = ""
                tags.append(("j_nodesc", "Benign by construction"))
            entry(post, post - dt.timedelta(days=R.randint(0, 4)), src, "System", d, two_lines(dr, cr, amt), tags)
    for m in range(1, 13):  # payroll, round totals by design
        post = weekday_on_or_after(dt.date(2025, m, 25))
        amt = 240000.00 + 1000 * R.randint(0, 25)
        entry(post, post, "SYS_PAY", "System", "Monthly payroll", [(6000, amt, 0.0), (1000, 0.0, amt)],
              [("j_rnd", "Benign by construction")])

    def manual(post, doc, user, desc, dr, cr, amt, tags):
        entry(post, doc, user, "Manual", desc, [(dr, amt, 0.0), (cr, 0.0, amt)], tags)

    for k, delta in enumerate([100.00, 1250.00, -75.50]):  # unbalanced
        post = rand_weekday(R)
        amt = amount()
        entry(post, post, R.choice(users_fin), "Manual", "Adjustment " + str(k + 1),
              [(6900, amt, 0.0), (2100, 0.0, round(amt + delta, 2))], [("j_unb", "Seeded anomaly")])
    for k in range(8):  # weekend manual
        post = rand_weekend(R)
        manual(post, post, R.choice(users_fin), "Weekend adjustment " + str(k + 1), 6900, 2100, amount(),
               [("j_wkd", "Seeded anomaly")])
    for k in range(6):  # late posting
        post = rand_weekday(R)
        post = weekday_on_or_after(min(post + dt.timedelta(days=60), dt.date(2025, 12, 26)) if post < dt.date(2025, 10, 30) else post)
        manual(post, post - dt.timedelta(days=R.randint(25, 60)), R.choice(users_fin), "Prior period correction " + str(k + 1),
               6900, 2100, amount(), [("j_late", "Seeded anomaly")])
    for k, u in enumerate(["TEMP02"] * 3 + ["ADMIN"] * 2):  # unauthorised users
        post = rand_weekday(R)
        manual(post, post, u, "Manual journal " + str(k + 1), 6900, 2100, amount(), [("j_unauth", "Seeded anomaly")])
    for k, (day, amt, desc) in enumerate([(29, 250000.00, "Year end sales adjustment"), (30, 180000.00, "Year end sales adjustment"),
                                          (30, 320000.00, ""), (31, 150000.00, "Sales cutoff entry"), (31, 410000.00, "")]):
        tags = [("j_mrev", "Seeded anomaly"), ("j_rnd", "Seeded anomaly"), ("j_pend", "Seeded anomaly")]
        if not desc:
            tags.append(("j_nodesc", "Seeded anomaly"))
        post = dt.date(2025, 12, day)
        manual(post, post, "FIN03", desc, 1000, 4000, amt, tags)
    for k in range(6):  # round manual
        post = rand_weekday(R)
        manual(post, post, R.choice(users_fin), "Round accrual " + str(k + 1), 6900, 2100, 1000.0 * R.randint(12, 45),
               [("j_rnd", "Seeded anomaly")])
    for k in range(7):  # missing description manual
        post = rand_weekday(R)
        manual(post, post, R.choice(users_fin), "", 6900, 2100, amount(), [("j_nodesc", "Seeded anomaly")])
    for k in range(5):  # below approval limit
        post = rand_weekday(R)
        manual(post, post, R.choice(users_fin), "Manual accrual " + str(k + 1), 6900, 2100, round(R.uniform(47600, 49950), 2),
               [("j_below", "Seeded anomaly")])
    for k in range(4):  # suspense
        post = rand_weekday(R)
        manual(post, post, R.choice(users_fin), "Suspense clearing " + str(k + 1), 9999, 1000, amount(),
               [("j_susp", "Seeded anomaly")])

    entries.sort(key=lambda e: (e["post"], R.random()))
    gap_idx = set(R.sample(range(60, len(entries) - 60), 6))
    n = 1000
    for i, e in enumerate(entries):
        n += 1
        if i in gap_idx:
            n += R.randint(1, 3)
            e["tags"].append(("j_gap", "Seeded anomaly"))
        e["no"] = n

    return dict(vendors=vendors, ven_bank=ven_bank, employees=employees, emp_bank=emp_bank, pays=pays, entries=entries)


# ----------------------------------------------------------------------------
# Independent pandas implementation of every test
# ----------------------------------------------------------------------------
BEN_EXP = [math.log10(1 + 1 / d) for d in range(1, 10)]


def first_digit(x):
    return int(f"{x:.2f}"[0])


def benford_stats(digits):
    n = len(digits)
    counts = [sum(1 for d in digits if d == k) for k in range(1, 10)]
    act = [c / n for c in counts]
    diff = [a - e for a, e in zip(act, BEN_EXP)]
    mad = sum(abs(x) for x in diff) / 9
    chi = sum(n * (a - e) ** 2 / e for a, e in zip(act, BEN_EXP))
    z = [max(0.0, abs(a - e) - 1 / (2 * n)) / math.sqrt(e * (1 - e) / n) for a, e in zip(act, BEN_EXP)]
    return dict(n=n, counts=counts, act=act, mad=mad, chi=chi, z=z)


def conformity(mad):
    if mad <= 0.006:
        return "Close conformity"
    if mad <= 0.012:
        return "Acceptable conformity"
    if mad <= 0.015:
        return "Marginally acceptable conformity"
    return "Nonconformity"


def pandas_tests(data):
    pv = PV
    p = pd.DataFrame(data["pays"])
    p["norm"] = p["inv"].str.upper().str.replace(r"[ /.]", "", regex=True)
    p["fd"] = p["amt"].map(first_digit)
    p["dup"] = p.duplicated(["vendor", "norm", "amt"], keep="first").astype(int)
    near = []
    for i, r in p.iterrows():
        m = p[(p.vendor == r.vendor) & (p.amt == r.amt) & ((p.date - r.date).map(lambda x: abs(x.days)) <= pv["dup_days"]) & (p.norm != r.norm)]
        near.append(int(len(m) > 0 and r.dup == 0))
    p["near"] = near
    cents = (p.amt * 100).round().astype(int)
    p["rnd"] = ((p.amt >= pv["round_min"]) & (cents % (pv["round_mod"] * 100) == 0)).astype(int)
    p["below"] = ((p.amt < pv["pay_limit"]) & (p.amt >= pv["pay_limit"] * (1 - pv["pay_band"]))).astype(int)
    g = p.groupby(["vendor", "date"]).amt.agg(["count", "sum"])
    p["split"] = [int(r.amt < pv["pay_limit"] and g.loc[(r.vendor, r.date), "count"] > 1 and g.loc[(r.vendor, r.date), "sum"] >= pv["pay_limit"])
                  for r in p.itertuples()]
    p["auth"] = ((p.amt >= pv["pay_limit"]) & (p.appr != pv["senior"])).astype(int)
    emp_banks = set(data["emp_bank"].values())
    vend_emp = {v: int(b in emp_banks) for v, b in data["ven_bank"].items()}
    p["emp"] = p.vendor.map(vend_emp)
    vb = {}
    for v in data["vendors"]:
        d = p[p.vendor == v].fd.tolist()
        n = len(d)
        if n:
            st = benford_stats(d)
            vb[v] = dict(n=n, mad=st["mad"], flag=int(n >= pv["ben_minn"] and st["mad"] > pv["ben_mad"]), emp=vend_emp[v])
        else:
            vb[v] = dict(n=0, mad=None, flag=0, emp=vend_emp[v])
    p["benf"] = p.vendor.map(lambda v: vb[v]["flag"])
    p["score"] = sum(p[k] * pv["wp_" + k] for k, _, _ in PAY_W)
    p["pri"] = (p.score >= pv["cut_pay"]).map({True: "High", False: ""})

    rows = []
    for e in data["entries"]:
        d = sum(l[1] for l in e["lines"])
        c = sum(l[2] for l in e["lines"])
        rows.append(dict(no=e["no"], post=e["post"], doc=e["doc"], user=e["user"], source=e["source"], desc=e["desc"],
                         debit=d, credit=c, rev=sum(1 for l in e["lines"] if CLASS[l[0]] == "Revenue"),
                         susp=sum(1 for l in e["lines"] if CLASS[l[0]] == "Suspense")))
    j = pd.DataFrame(rows).sort_values("no").reset_index(drop=True)
    j["unb"] = ((j.debit - j.credit).abs() > pv["tol"]).astype(int)
    j["wkd"] = j.post.map(lambda x: int(x.weekday() >= 5))
    j["late"] = ((j.post - j.doc).map(lambda x: x.days) > pv["je_lag"]).astype(int)
    j["pend"] = j.post.map(lambda x: int(pv["fy_end"] - dt.timedelta(days=pv["je_end"]) < x <= pv["fy_end"]))
    j["unauth"] = (~j.user.isin(AUTH_USERS)).astype(int)
    j["mrev"] = ((j.source == "Manual") & (j.rev > 0)).astype(int)
    jc = (j.debit * 100).round().astype(int)
    j["rnd"] = ((j.debit >= pv["je_round_min"]) & (jc % (pv["je_round_mod"] * 100) == 0)).astype(int)
    j["nodesc"] = (j.desc.str.strip() == "").astype(int)
    j["below"] = ((j.source == "Manual") & (j.debit < pv["je_limit"]) & (j.debit >= pv["je_limit"] * (1 - pv["je_band"]))).astype(int)
    j["susp_f"] = ((j.source == "Manual") & (j.susp > 0)).astype(int)
    j["gap"] = (j.no.diff().fillna(1) > 1).astype(int)
    j["score"] = sum(j[k if k != "susp" else "susp_f"] * pv["wj_" + k] for k, _, _ in JE_W)
    j["pri"] = (j.score >= pv["cut_je"]).map({True: "High", False: ""})

    ben = benford_stats(p.fd.tolist())
    return p, j, ben, vb


def build_answer_key(data, p, j, vb):
    ak = []
    for r in p.itertuples():
        for t, cat in r.tags:
            ak.append(("Payment", r.id, t, cat, r.note))
        if r.note in ("Monthly rent", "Monthly licence") and r.near == 1:
            ak.append(("Payment", r.id, "p_near", "Benign by construction", r.note + ", same amount within the window"))
        if r.vendor == "V031" and r.split == 1:
            ak.append(("Payment", r.id, "p_split", "Counterpart of seeded anomaly", "Same day collision of threshold avoidance payments"))
    jflag = {t: t[2:] if t != "j_susp" else "susp_f" for t in ["j_" + k for k, _, _ in JE_W]}
    jtags = {e["no"]: e["tags"] for e in data["entries"]}
    for no in j.no:
        for t, cat in jtags[no]:
            ak.append(("Journal entry", int(no), t, cat, ""))
    ak.append(("Vendor", "V031", "v_benf", "Seeded anomaly", "Payments clustered below the approval limit"))
    ak.append(("Vendor", "V038", "v_emp", "Seeded anomaly", "Vendor bank account equals employee E014 bank account"))
    return ak


# ----------------------------------------------------------------------------
# Workbook
# ----------------------------------------------------------------------------
HDR_FILL = PatternFill("solid", fgColor="1F3864")
HDR_FONT = Font(bold=True, color="FFFFFF")
INPUT_FONT = Font(color="0000FF")
ENGINE_FONT = Font(color="7F6000")
TITLE_FONT = Font(bold=True, size=14)
NOTE_FONT = Font(italic=True, color="595959")
BOLD = Font(bold=True)
MONEY = '#,##0.00;(#,##0.00)'
DATEF = 'dd mmm yyyy'
PCT = '0.0%'


def header(ws, row, labels, widths=None):
    for i, lab in enumerate(labels, start=1):
        c = ws.cell(row=row, column=i, value=lab)
        c.fill, c.font = HDR_FILL, HDR_FONT
        c.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(i)].width = (widths or {}).get(i, 14)
    ws.row_dimensions[row].height = 34


def title(ws, text, note=None):
    ws["A1"] = text
    ws["A1"].font = TITLE_FONT
    ws["A2"] = SAMPLE_LABEL
    ws["A2"].font = NOTE_FONT
    if note:
        ws["A3"] = note
        ws["A3"].font = NOTE_FONT


def build_workbook(data, p, j, ben, vb, ak, path):
    wb = Workbook()
    nP, nJ = len(p), len(j)
    nLines = sum(len(e["lines"]) for e in data["entries"])
    pf, pl = 5, 4 + nP           # payments first and last row
    jf, jl = 5, 4 + nJ           # JE entries first and last row
    lf, ll = 5, 4 + nLines       # JE lines
    nV = len(data["vendors"])
    vf, vl = 5, 4 + nV           # Vendors master
    vbf, vbl = 6, 5 + nV         # Vendor_Benford
    nE = len(data["employees"])
    ef, el = 5, 4 + nE

    # ---------------- README ----------------
    ws = wb.active
    ws.title = "README"
    ws.column_dimensions["A"].width = 125
    lines = [
        ("Audit analytics toolkit: Benford, duplicate payments, round sums and journal entry tests", TITLE_FONT),
        (SAMPLE_LABEL, NOTE_FONT), ("", None),
        ("PURPOSE", BOLD),
        ("Run standard data analytics over a year of payments and journal entries, score each item, and produce a ranked list for targeted follow up. "
         "Anomalies were seeded into the sample data so detection can be tested against an answer key.", None), ("", None),
        ("SHEETS", BOLD),
        ("Parameters: thresholds, weights, authorised users (inputs, blue). Change a value and every flag recalculates.", None),
        ("Vendors, Employees: master data used by the bank account match.", None),
        ("Payments: one row per payment. Blue = input, black = formula. Flags and score are columns to the right.", None),
        ("Benford: first digit test over all payments (counts, z statistics, MAD, chi square) with a chart.", None),
        ("Vendor_Benford: the same test per vendor, plus the employee bank account match.", None),
        ("COA, JE_Lines: chart of accounts and journal entry lines (inputs).", None),
        ("JE_Entries: one row per journal entry number (number in brown from the engine, everything else formulas), with tests and score.", None),
        ("Exceptions: counts per test, recall against the answer key, and ranked High risk lists.", None),
        ("Answer_Key: seeded anomalies and known benign patterns, with a live Detected? column.", None),
        ("Checks: control totals and completeness. Overall result must read PASS.", None), ("", None),
        ("HOW TO READ THE RESULTS", BOLD),
        ("A flag is a reason to look, not a finding. Each High item needs a documented disposition by the reviewer.", None),
        ("The High risk list is a targeted selection of specific items. It is not an audit sample and results must not be projected to the population (ISA 530).", None),
        ("Benford conformity is a screening result. Nonconformity can come from natural causes such as fixed price contracts or recurring amounts.", None), ("", None),
        ("ASSUMPTIONS AND LIMITS", BOLD),
        ("1. Fiscal year 1 Jan 2025 to 31 Dec 2025. All data is synthetic, generated with a fixed seed.", None),
        ("2. Benford first digit test uses Nigrini style MAD bands: up to 0.006 close, to 0.012 acceptable, to 0.015 marginal, above nonconformity.", None),
        ("3. Vendor Benford results are ignored below the minimum payment count on the Parameters sheet.", None),
        ("4. Duplicate test ignores spaces, slashes, dots and letter case in invoice numbers. It does not catch transposed digits.", None),
        ("5. Journal entry user and source are taken from the first line of each entry.", None),
        ("6. Revised ISA 240 applies to audits of periods beginning on or after 15 Dec 2026. Use it as a reference for fraud risk work, not as a current requirement for earlier periods.", None),
        ("7. Real data would need data integrity checks (completeness and accuracy of the extract) before any of these tests are run.", None), ("", None),
        ("COLOUR CODE", BOLD),
        ("Blue font = input. Brown font = engine output pasted as values. Black = formula.", None),
    ]
    for i, (t, f) in enumerate(lines, start=1):
        c = ws.cell(row=i, column=1, value=t)
        c.alignment = Alignment(wrap_text=True)
        if f:
            c.font = f

    # ---------------- Parameters ----------------
    ws = wb.create_sheet("Parameters")
    title(ws, "Parameters", "Inputs in blue. Weights drive the risk scores.")
    header(ws, 4, ["Parameter", "Value", "Note"], {1: 50, 2: 16, 3: 70})
    PA = {}
    r = 5
    for k, lab, val, note in PARAMS:
        ws.cell(row=r, column=1, value=lab)
        c = ws.cell(row=r, column=2, value=val)
        c.font = INPUT_FONT
        if isinstance(val, dt.date):
            c.number_format = DATEF
        ws.cell(row=r, column=3, value=note)
        PA[k] = f"Parameters!$B${r}"
        r += 1
    r += 1
    ws.cell(row=r, column=1, value="Payment risk weights").font = BOLD
    r += 1
    for k, lab, w in PAY_W:
        ws.cell(row=r, column=1, value=lab)
        ws.cell(row=r, column=2, value=w).font = INPUT_FONT
        PA["wp_" + k] = f"Parameters!$B${r}"
        r += 1
    r += 1
    ws.cell(row=r, column=1, value="Journal entry risk weights").font = BOLD
    r += 1
    for k, lab, w in JE_W:
        ws.cell(row=r, column=1, value=lab)
        ws.cell(row=r, column=2, value=w).font = INPUT_FONT
        PA["wj_" + k] = f"Parameters!$B${r}"
        r += 1
    r += 1
    ws.cell(row=r, column=1, value="Authorised journal posting users").font = BOLD
    r += 1
    u0 = r
    for u in AUTH_USERS:
        ws.cell(row=r, column=1, value=u).font = INPUT_FONT
        r += 1
    AUTHR = f"Parameters!$A${u0}:$A${r - 1}"

    # ---------------- Vendors, Employees ----------------
    ws = wb.create_sheet("Vendors")
    title(ws, "Vendor master", "Synthetic. V038 bank account equals an employee bank account (seeded).")
    header(ws, 4, ["Vendor ID", "Vendor name", "Bank account"], {2: 26, 3: 22})
    for i, v in enumerate(data["vendors"]):
        for c, val in enumerate([v, f"Sample Supplier {v[1:]}", data["ven_bank"][v]], start=1):
            ws.cell(row=5 + i, column=c, value=val).font = INPUT_FONT
    ws = wb.create_sheet("Employees")
    title(ws, "Employee master", "Synthetic.")
    header(ws, 4, ["Employee ID", "Bank account"], {2: 22})
    for i, e in enumerate(data["employees"]):
        for c, val in enumerate([e, data["emp_bank"][e]], start=1):
            ws.cell(row=5 + i, column=c, value=val).font = INPUT_FONT

    # ---------------- Payments ----------------
    PCOLS = ["id", "date", "vendor", "inv", "amt", "appr", "fd", "norm", "dup", "near", "rnd", "below", "split", "auth", "emp", "benf", "score", "pri"]
    PH = ["Payment ID", "Payment date", "Vendor", "Invoice no", "Amount (USD)", "Approver", "First digit", "Normalised invoice",
          "Duplicate", "Near duplicate", "Round sum", "Below limit", "Same day split", "Authority breach", "Employee bank match",
          "Vendor Benford flag", "Risk score", "Priority"]
    PL = {k: get_column_letter(i + 1) for i, k in enumerate(PCOLS)}

    def prng(k):
        return f"${PL[k]}${pf}:${PL[k]}${pl}"

    ws = wb.create_sheet("Payments")
    title(ws, "Payments", "Columns A to F are inputs. Columns G onward are formulas.")
    header(ws, 4, PH, {1: 11, 2: 13, 4: 14, 5: 14, 8: 18, 9: 11, 10: 12, 15: 14, 16: 14})
    for i, rr in enumerate(p.itertuples()):
        r = pf + i
        for c, val in enumerate([rr.id, rr.date, rr.vendor, rr.inv, rr.amt, rr.appr], start=1):
            ws.cell(row=r, column=c, value=val).font = INPUT_FONT
        L = PL
        ws[f"{L['fd']}{r}"] = f'=VALUE(LEFT(TEXT({L["amt"]}{r},"0.00"),1))'
        ws[f"{L['norm']}{r}"] = f'=UPPER(SUBSTITUTE(SUBSTITUTE(SUBSTITUTE({L["inv"]}{r}," ",""),"/",""),".",""))'
        ws[f"{L['dup']}{r}"] = (f'=IF(COUNTIFS(${L["vendor"]}${pf}:{L["vendor"]}{r},{L["vendor"]}{r},${L["norm"]}${pf}:{L["norm"]}{r},{L["norm"]}{r},'
                                f'${L["amt"]}${pf}:{L["amt"]}{r},{L["amt"]}{r})>1,1,0)')
        ws[f"{L['near']}{r}"] = (f'=IF(AND({L["dup"]}{r}=0,COUNTIFS({prng("vendor")},{L["vendor"]}{r},{prng("amt")},{L["amt"]}{r},'
                                 f'{prng("date")},">="&({L["date"]}{r}-{PA["dup_days"]}),{prng("date")},"<="&({L["date"]}{r}+{PA["dup_days"]}),'
                                 f'{prng("norm")},"<>"&{L["norm"]}{r})>0),1,0)')
        ws[f"{L['rnd']}{r}"] = f'=IF(AND({L["amt"]}{r}>={PA["round_min"]},ROUND(MOD(ROUND({L["amt"]}{r},2),{PA["round_mod"]}),2)=0),1,0)'
        ws[f"{L['below']}{r}"] = f'=IF(AND({L["amt"]}{r}<{PA["pay_limit"]},{L["amt"]}{r}>={PA["pay_limit"]}*(1-{PA["pay_band"]})),1,0)'
        ws[f"{L['split']}{r}"] = (f'=IF(AND({L["amt"]}{r}<{PA["pay_limit"]},COUNTIFS({prng("vendor")},{L["vendor"]}{r},{prng("date")},{L["date"]}{r})>1,'
                                  f'SUMIFS({prng("amt")},{prng("vendor")},{L["vendor"]}{r},{prng("date")},{L["date"]}{r})>={PA["pay_limit"]}),1,0)')
        ws[f"{L['auth']}{r}"] = f'=IF(AND({L["amt"]}{r}>={PA["pay_limit"]},{L["appr"]}{r}<>{PA["senior"]}),1,0)'
        ws[f"{L['emp']}{r}"] = f'=INDEX(Vendor_Benford!$N${vbf}:$N${vbl},MATCH({L["vendor"]}{r},Vendor_Benford!$A${vbf}:$A${vbl},0))'
        ws[f"{L['benf']}{r}"] = f'=INDEX(Vendor_Benford!$M${vbf}:$M${vbl},MATCH({L["vendor"]}{r},Vendor_Benford!$A${vbf}:$A${vbl},0))'
        ws[f"{L['score']}{r}"] = "=" + "+".join(f'{L[k]}{r}*{PA["wp_" + k]}' for k, _, _ in PAY_W)
        ws[f"{L['pri']}{r}"] = f'=IF({L["score"]}{r}>={PA["cut_pay"]},"High","")'
        ws.cell(row=r, column=2).number_format = DATEF
        ws.cell(row=r, column=5).number_format = MONEY
    ws[f"{PL['dup']}4"].comment = Comment("1 for the later occurrence of the same vendor, normalised invoice number and amount.", "Preparer")
    ws[f"{PL['near']}4"].comment = Comment("1 if another payment has the same vendor and amount within the window but a different invoice number. Flags both payments.", "Preparer")
    ws[f"{PL['below']}4"].comment = Comment("Amount is below the approval limit and within the band under it.", "Preparer")
    ws[f"{PL['split']}4"].comment = Comment("Payment is below the limit, but payments to the same vendor on the same date add up to the limit or more.", "Preparer")
    ws.freeze_panes = "A5"

    # ---------------- Benford (all payments) ----------------
    ws = wb.create_sheet("Benford")
    title(ws, "Benford's Law, first digit test, all payments", "z uses the continuity correction. MAD is the mean absolute deviation across nine digits.")
    header(ws, 4, ["First digit", "Count", "Actual %", "Expected %", "Difference", "z statistic", "Significant at 5 percent?", "Chi square term"],
           {1: 11, 7: 22, 8: 16})
    fdr = f"Payments!${PL['fd']}${pf}:${PL['fd']}${pl}"
    for d in range(1, 10):
        r = 4 + d
        ws.cell(row=r, column=1, value=d)
        ws.cell(row=r, column=2, value=f"=COUNTIF({fdr},A{r})")
        ws.cell(row=r, column=3, value=f"=B{r}/$B$15")
        ws.cell(row=r, column=4, value=f"=LOG10(1+1/A{r})")
        ws.cell(row=r, column=5, value=f"=C{r}-D{r}")
        ws.cell(row=r, column=6, value=f"=MAX(0,ABS(E{r})-1/(2*$B$15))/SQRT(D{r}*(1-D{r})/$B$15)")
        ws.cell(row=r, column=7, value=f'=IF(F{r}>{PA["ben_z"]},"Yes","No")')
        ws.cell(row=r, column=8, value=f"=$B$15*(C{r}-D{r})^2/D{r}")
        for c in (3, 4, 5):
            ws.cell(row=r, column=c).number_format = '0.00%'
        ws.cell(row=r, column=6).number_format = '0.00'
        ws.cell(row=r, column=8).number_format = '0.00'
    ws["A15"], ws["B15"] = "Total payments (n)", "=SUM(B5:B13)"
    ws["A16"], ws["B16"] = "MAD", "=SUMPRODUCT(ABS(E5:E13))/9"
    ws["A17"], ws["B17"] = "Conformity (MAD bands)", '=IF(B16<=0.006,"Close conformity",IF(B16<=0.012,"Acceptable conformity",IF(B16<=0.015,"Marginally acceptable conformity","Nonconformity")))'
    ws["A18"], ws["B18"] = "Chi square", "=SUM(H5:H13)"
    ws["A19"], ws["B19"] = "Chi square critical (8 df, 5 percent)", f"={PA['ben_chi']}"
    ws["A20"], ws["B20"] = "Chi square result", '=IF(B18>B19,"Reject conformity","Do not reject")'
    ws["A21"], ws["B21"] = "Digits significant at 5 percent", '=COUNTIF(G5:G13,"Yes")'
    ws["B16"].number_format = '0.0000'
    ws["B18"].number_format = '0.00'
    ws.column_dimensions["A"].width = 36
    ws["B16"].comment = Comment("Mean absolute deviation of actual from expected first digit proportions.", "Preparer")
    ch = BarChart()
    ch.type, ch.title = "col", "First digit: actual vs expected (sample data)"
    ch.add_data(Reference(ws, min_col=3, min_row=4, max_row=13), titles_from_data=True)
    ch.add_data(Reference(ws, min_col=4, min_row=4, max_row=13), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=1, min_row=5, max_row=13))
    ch.height, ch.width = 8, 16
    ws.add_chart(ch, "D15")

    # ---------------- Vendor_Benford ----------------
    ws = wb.create_sheet("Vendor_Benford")
    title(ws, "Vendor level Benford test and employee bank match", "Expected proportions in row 5. Flag needs the minimum payment count on the Parameters sheet.")
    header(ws, 4, ["Vendor", "Payments (n)"] + [f"Digit {d}" for d in range(1, 10)] + ["MAD", "Benford flag", "Employee bank match"],
           {1: 11, 12: 11, 13: 13, 14: 16})
    ws.cell(row=5, column=1, value="Expected")
    for d in range(1, 10):
        c = ws.cell(row=5, column=2 + d, value=f"=LOG10(1+1/{d})")
        c.number_format = '0.0000'
    for i, v in enumerate(data["vendors"]):
        r = vbf + i
        ws.cell(row=r, column=1, value=v)
        ws.cell(row=r, column=2, value=f"=COUNTIF(Payments!{prng('vendor')},A{r})")
        for d in range(1, 10):
            ws.cell(row=r, column=2 + d, value=f"=COUNTIFS(Payments!{prng('vendor')},$A{r},{fdr},{d})")
        ws.cell(row=r, column=12, value=f'=IF(B{r}=0,0,SUMPRODUCT(ABS(C{r}:K{r}/B{r}-$C$5:$K$5))/9)')
        ws.cell(row=r, column=13, value=f'=IF(AND(B{r}>={PA["ben_minn"]},L{r}>{PA["ben_mad"]}),1,0)')
        ws.cell(row=r, column=14, value=f'=IF(COUNTIF(Employees!$B${ef}:$B${el},INDEX(Vendors!$C${vf}:$C${vl},MATCH(A{r},Vendors!$A${vf}:$A${vl},0)))>0,1,0)')
        ws.cell(row=r, column=12).number_format = '0.0000'
    ws.freeze_panes = "B6"

    # ---------------- COA, JE_Lines ----------------
    ws = wb.create_sheet("COA")
    title(ws, "Chart of accounts", "Synthetic.")
    header(ws, 4, ["Account", "Name", "Class"], {2: 24})
    for i, (a, nme, cl) in enumerate(COA):
        for c, val in enumerate([a, nme, cl], start=1):
            ws.cell(row=5 + i, column=c, value=val).font = INPUT_FONT
    coa_last = 4 + len(COA)
    ws = wb.create_sheet("JE_Lines")
    title(ws, "Journal entry lines", "Inputs in blue. Class is a formula.")
    header(ws, 4, ["Line ID", "JE no", "Posting date", "Document date", "Account", "Class", "Debit (USD)", "Credit (USD)", "User", "Source", "Description"],
           {3: 14, 4: 14, 7: 14, 8: 14, 11: 34})
    lr = lf
    for e in sorted(data["entries"], key=lambda x: x["no"]):
        for (acct, dr, cr) in e["lines"]:
            vals = [f"L{lr - lf + 1:05d}", e["no"], e["post"], e["doc"], acct, None, dr, cr, e["user"], e["source"], e["desc"] or None]
            for c, val in enumerate(vals, start=1):
                if c == 6:
                    ws.cell(row=lr, column=6, value=f"=INDEX(COA!$C$5:$C${coa_last},MATCH(E{lr},COA!$A$5:$A${coa_last},0))")
                elif val is not None:
                    ws.cell(row=lr, column=c, value=val).font = INPUT_FONT
            ws.cell(row=lr, column=3).number_format = DATEF
            ws.cell(row=lr, column=4).number_format = DATEF
            ws.cell(row=lr, column=7).number_format = MONEY
            ws.cell(row=lr, column=8).number_format = MONEY
            lr += 1
    assert lr - 1 == ll
    ws.freeze_panes = "A5"

    def lrng(col):
        return f"JE_Lines!${col}${lf}:${col}${ll}"

    # ---------------- JE_Entries ----------------
    JC = ["no", "n", "debit", "credit", "diff", "post", "doc", "user", "source", "desc", "lag", "rev", "susp_n",
          "unb", "wkd", "late", "pend", "unauth", "mrev", "rnd", "nodesc", "below", "susp", "gap", "score", "pri"]
    JH = ["JE no", "Lines", "Total debit", "Total credit", "Debit less credit", "Posting date", "Document date", "User", "Source", "Description",
          "Lag (days)", "Revenue lines", "Suspense lines", "Unbalanced", "Weekend", "Late posting", "Year end window", "Unauthorised user",
          "Manual to revenue", "Round total", "No description", "Below limit", "Manual to suspense", "Sequence gap", "Risk score", "Priority"]
    JL = {k: get_column_letter(i + 1) for i, k in enumerate(JC)}
    ws = wb.create_sheet("JE_Entries")
    title(ws, "Journal entry tests, one row per entry", "JE number is from the engine (brown). All other columns are formulas.")
    header(ws, 4, JH, {1: 10, 3: 14, 4: 14, 5: 13, 6: 13, 7: 13, 10: 28, 14: 11, 18: 14, 19: 14, 24: 11})
    for i, rr in enumerate(j.itertuples()):
        r = jf + i
        L = JL
        m = f"MATCH({L['no']}{r},{lrng('B')},0)"
        ws[f"{L['no']}{r}"] = int(rr.no)
        ws[f"{L['no']}{r}"].font = ENGINE_FONT
        ws[f"{L['n']}{r}"] = f"=COUNTIF({lrng('B')},{L['no']}{r})"
        ws[f"{L['debit']}{r}"] = f"=SUMIFS({lrng('G')},{lrng('B')},{L['no']}{r})"
        ws[f"{L['credit']}{r}"] = f"=SUMIFS({lrng('H')},{lrng('B')},{L['no']}{r})"
        ws[f"{L['diff']}{r}"] = f"={L['debit']}{r}-{L['credit']}{r}"
        ws[f"{L['post']}{r}"] = f"=INDEX({lrng('C')},{m})"
        ws[f"{L['doc']}{r}"] = f"=INDEX({lrng('D')},{m})"
        ws[f"{L['user']}{r}"] = f"=INDEX({lrng('I')},{m})"
        ws[f"{L['source']}{r}"] = f"=INDEX({lrng('J')},{m})"
        ws[f"{L['desc']}{r}"] = f'=IF(LEN(INDEX({lrng("K")},{m}))=0,"",INDEX({lrng("K")},{m}))'
        ws[f"{L['lag']}{r}"] = f"={L['post']}{r}-{L['doc']}{r}"
        ws[f"{L['rev']}{r}"] = f'=COUNTIFS({lrng("B")},{L["no"]}{r},{lrng("F")},"Revenue")'
        ws[f"{L['susp_n']}{r}"] = f'=COUNTIFS({lrng("B")},{L["no"]}{r},{lrng("F")},"Suspense")'
        ws[f"{L['unb']}{r}"] = f'=IF(ABS({L["diff"]}{r})>{PA["tol"]},1,0)'
        ws[f"{L['wkd']}{r}"] = f"=IF(WEEKDAY({L['post']}{r},2)>5,1,0)"
        ws[f"{L['late']}{r}"] = f'=IF({L["lag"]}{r}>{PA["je_lag"]},1,0)'
        ws[f"{L['pend']}{r}"] = f'=IF(AND({L["post"]}{r}<={PA["fy_end"]},{L["post"]}{r}>{PA["fy_end"]}-{PA["je_end"]}),1,0)'
        ws[f"{L['unauth']}{r}"] = f"=IF(COUNTIF({AUTHR},{L['user']}{r})=0,1,0)"
        ws[f"{L['mrev']}{r}"] = f'=IF(AND({L["source"]}{r}="Manual",{L["rev"]}{r}>0),1,0)'
        ws[f"{L['rnd']}{r}"] = f'=IF(AND({L["debit"]}{r}>={PA["je_round_min"]},ROUND(MOD(ROUND({L["debit"]}{r},2),{PA["je_round_mod"]}),2)=0),1,0)'
        ws[f"{L['nodesc']}{r}"] = f'=IF(LEN(TRIM({L["desc"]}{r}))=0,1,0)'
        ws[f"{L['below']}{r}"] = (f'=IF(AND({L["source"]}{r}="Manual",{L["debit"]}{r}<{PA["je_limit"]},'
                                  f'{L["debit"]}{r}>={PA["je_limit"]}*(1-{PA["je_band"]})),1,0)')
        ws[f"{L['susp']}{r}"] = f'=IF(AND({L["source"]}{r}="Manual",{L["susp_n"]}{r}>0),1,0)'
        ws[f"{L['gap']}{r}"] = "=0" if i == 0 else f"=IF({L['no']}{r}-{L['no']}{r - 1}>1,1,0)"
        ws[f"{L['score']}{r}"] = "=" + "+".join(f'{L[k]}{r}*{PA["wj_" + k]}' for k, _, _ in JE_W)
        ws[f"{L['pri']}{r}"] = f'=IF({L["score"]}{r}>={PA["cut_je"]},"High","")'
        for k in ("debit", "credit", "diff"):
            ws[f"{L[k]}{r}"].number_format = MONEY
        for k in ("post", "doc"):
            ws[f"{L[k]}{r}"].number_format = DATEF
    ws[f"{JL['pend']}4"].comment = Comment("Posting date falls in the last days of the fiscal year (window on Parameters).", "Preparer")
    ws[f"{JL['gap']}4"].comment = Comment("1 if the JE number jumps from the previous entry, meaning one or more numbers are missing.", "Preparer")
    ws.freeze_panes = "B5"

    def jr(k):
        return f"JE_Entries!${JL[k]}${jf}:${JL[k]}${jl}"

    # ---------------- Answer_Key (rows first so counts can reference) ----------------
    ak_rows = []
    pflag = {"p_dup": "dup", "p_near": "near", "p_rnd": "rnd", "p_below": "below", "p_split": "split", "p_auth": "auth"}
    pmap = p.set_index("id")
    jmap = j.set_index("no")
    for lvl, ident, t, cat, note in ak:
        ak_rows.append((lvl, ident, t, cat, note))
    # add benign by construction rows for the payments tags already carried (they are in ak)
    wsA = wb.create_sheet("Answer_Key")
    title(wsA, "Answer key: seeded anomalies and known benign patterns",
          "For validation only. Detected? is a live lookup into the flag columns. Benign rows are expected false positives.")
    header(wsA, 4, ["Level", "ID", "Test key", "Category", "Note", "Detected?"], {1: 14, 2: 12, 3: 12, 4: 32, 5: 52, 6: 11})
    for i, (lvl, ident, t, cat, note) in enumerate(ak_rows):
        r = 5 + i
        for c, val in enumerate([lvl, ident, t, cat, note], start=1):
            wsA.cell(row=r, column=c, value=val).font = ENGINE_FONT
        if lvl == "Payment":
            col = PL[pflag[t]]
            f = f'=IF(INDEX(Payments!${col}${pf}:${col}${pl},MATCH(B{r},Payments!$A${pf}:$A${pl},0))=1,"Yes","No")'
        elif lvl == "Journal entry":
            col = JL[t[2:] if t != "j_susp" else "susp"]
            f = f'=IF(INDEX(JE_Entries!${col}${jf}:${col}${jl},MATCH(B{r},JE_Entries!$A${jf}:$A${jl},0))=1,"Yes","No")'
        else:
            col = "M" if t == "v_benf" else "N"
            f = f'=IF(INDEX(Vendor_Benford!${col}${vbf}:${col}${vbl},MATCH(B{r},Vendor_Benford!$A${vbf}:$A${vbl},0))=1,"Yes","No")'
        wsA.cell(row=r, column=6, value=f)
    akl = 4 + len(ak_rows)
    AK = lambda col: f"Answer_Key!${col}$5:${col}${akl}"

    # ---------------- Exceptions ----------------
    ws = wb.create_sheet("Exceptions")
    title(ws, "Exceptions summary and High risk lists", "Recall is measured against the answer key. Flagged but not in key = items a reviewer must clear.")
    header(ws, 4, ["Test key", "Test", "Population", "Flagged", "Flag rate", "Seeded anomalies", "Seeded detected", "Recall", "Flagged, not in answer key"],
           {1: 11, 2: 56, 3: 12, 4: 10, 5: 10, 6: 12, 7: 12, 8: 9, 9: 18})
    tests = [("p_dup", "Duplicate payment (vendor, normalised invoice, amount)", "Payments", PL["dup"]),
             ("p_near", "Near duplicate payment", "Payments", PL["near"]),
             ("p_rnd", "Round sum payment", "Payments", PL["rnd"]),
             ("p_below", "Payment just below approval limit", "Payments", PL["below"]),
             ("p_split", "Same day split payments above limit", "Payments", PL["split"]),
             ("p_auth", "Approved below required authority", "Payments", PL["auth"]),
             ("v_benf", "Vendor fails Benford test", "Vendor_Benford", "M"),
             ("v_emp", "Vendor bank account matches employee", "Vendor_Benford", "N")]
    jt = [("j_unb", "unb", "Unbalanced journal entry"), ("j_wkd", "wkd", "Posted on a weekend"), ("j_late", "late", "Late posting versus document date"),
          ("j_pend", "pend", "Posted in last days of the year"), ("j_unauth", "unauth", "Posted by unauthorised user"),
          ("j_mrev", "mrev", "Manual entry to revenue"), ("j_rnd", "rnd", "Round total entry"), ("j_nodesc", "nodesc", "Entry without description"),
          ("j_below", "below", "Manual entry just below approval limit"), ("j_susp", "susp", "Manual entry to suspense"), ("j_gap", "gap", "Sequence gap before entry")]
    for k, lab, kk in [(a, c, b) for a, b, c in jt]:
        tests.append((k, lab, "JE_Entries", JL[kk]))
    EXC = {}
    for i, (key, lab, sheet, col) in enumerate(tests):
        r = 5 + i
        EXC[key] = r
        if sheet == "Payments":
            rng, pop = f"Payments!${col}${pf}:${col}${pl}", f"=COUNTA(Payments!$A${pf}:$A${pl})"
        elif sheet == "Vendor_Benford":
            rng, pop = f"Vendor_Benford!${col}${vbf}:${col}${vbl}", f"=COUNTA(Vendor_Benford!$A${vbf}:$A${vbl})"
        else:
            rng, pop = f"JE_Entries!${col}${jf}:${col}${jl}", f"=COUNTA(JE_Entries!$A${jf}:$A${jl})"
        ws.cell(row=r, column=1, value=key)
        ws.cell(row=r, column=2, value=lab)
        ws.cell(row=r, column=3, value=pop)
        ws.cell(row=r, column=4, value=f"=SUM({rng})")
        ws.cell(row=r, column=5, value=f"=D{r}/C{r}").number_format = PCT
        ws.cell(row=r, column=6, value=f'=COUNTIFS({AK("C")},A{r},{AK("D")},"Seeded anomaly")')
        ws.cell(row=r, column=7, value=f'=COUNTIFS({AK("C")},A{r},{AK("D")},"Seeded anomaly",{AK("F")},"Yes")')
        ws.cell(row=r, column=8, value=f'=IF(F{r}=0,"n/a",G{r}/F{r})').number_format = '0%'
        ws.cell(row=r, column=9, value=f'=D{r}-COUNTIFS({AK("C")},A{r},{AK("F")},"Yes")')
    ws["I4"].comment = Comment("Flagged items that are neither seeded, counterparts, nor known benign patterns. A reviewer would have to clear these in practice.", "Preparer")
    tl = 4 + len(tests)

    # High risk lists (IDs from engine, details from formulas)
    hp = p[p.pri == "High"].sort_values(["score", "date"], ascending=[False, True])
    hj = j[j.pri == "High"].sort_values(["score", "post"], ascending=[False, True])
    r0 = tl + 3
    ws.cell(row=r0 - 1, column=1, value="High risk payments (ranked by engine, details are formulas)").font = BOLD
    for c, h in enumerate(["Payment ID", "Vendor", "Date", "Amount (USD)", "Score", "Reviewer disposition (input)"], start=1):
        cc = ws.cell(row=r0, column=c, value=h)
        cc.fill, cc.font = HDR_FILL, HDR_FONT
    for i, pid in enumerate(hp.id):
        r = r0 + 1 + i
        m = f"MATCH($A{r},Payments!$A${pf}:$A${pl},0)"
        ws.cell(row=r, column=1, value=pid).font = ENGINE_FONT
        ws.cell(row=r, column=2, value=f"=INDEX(Payments!$C${pf}:$C${pl},{m})")
        ws.cell(row=r, column=3, value=f"=INDEX(Payments!$B${pf}:$B${pl},{m})").number_format = DATEF
        ws.cell(row=r, column=4, value=f"=INDEX(Payments!$E${pf}:$E${pl},{m})").number_format = MONEY
        ws.cell(row=r, column=5, value=f"=INDEX(Payments!${PL['score']}${pf}:${PL['score']}${pl},{m})")
        ws.cell(row=r, column=6).font = INPUT_FONT
    pend_row = r0 + len(hp)
    r1 = pend_row + 4
    ws.cell(row=r1 - 1, column=1, value="High risk journal entries (ranked by engine, details are formulas)").font = BOLD
    for c, h in enumerate(["JE no", "User", "Posting date", "Total debit (USD)", "Score", "Reviewer disposition (input)"], start=1):
        cc = ws.cell(row=r1, column=c, value=h)
        cc.fill, cc.font = HDR_FILL, HDR_FONT
    for i, no in enumerate(hj.no):
        r = r1 + 1 + i
        m = f"MATCH($A{r},JE_Entries!$A${jf}:$A${jl},0)"
        ws.cell(row=r, column=1, value=int(no)).font = ENGINE_FONT
        ws.cell(row=r, column=2, value=f"=INDEX(JE_Entries!${JL['user']}${jf}:${JL['user']}${jl},{m})")
        ws.cell(row=r, column=3, value=f"=INDEX(JE_Entries!${JL['post']}${jf}:${JL['post']}${jl},{m})").number_format = DATEF
        ws.cell(row=r, column=4, value=f"=INDEX(JE_Entries!${JL['debit']}${jf}:${JL['debit']}${jl},{m})").number_format = MONEY
        ws.cell(row=r, column=5, value=f"=INDEX(JE_Entries!${JL['score']}${jf}:${JL['score']}${jl},{m})")
        ws.cell(row=r, column=6).font = INPUT_FONT
    jend_row = r1 + len(hj)
    ws.column_dimensions["F"].width = 28

    # ---------------- Checks ----------------
    ws = wb.create_sheet("Checks")
    title(ws, "Controls", "Differences must be zero. Overall result must read PASS.")
    header(ws, 4, ["Control", "Result", "Expected", "Pass?"], {1: 70, 2: 16, 3: 16, 4: 10})
    ck = [
        ("Payments listed less Benford digit counts", f"=COUNTA(Payments!$A${pf}:$A${pl})-Benford!B15", 0),
        ("Payments listed less Vendor_Benford counts", f"=COUNTA(Payments!$A${pf}:$A${pl})-SUM(Vendor_Benford!$B${vbf}:$B${vbl})", 0),
        ("JE lines in JE_Lines less lines assigned to entries", f"=COUNTA({lrng('A')})-SUM({jr('n')})", 0),
        ("Total debits less total credits per JE_Lines, less sum of entry differences", f"=SUM({lrng('G')})-SUM({lrng('H')})-SUM({jr('diff')})", 0),
        ("Every JE line has a class (count of errors)", f"=SUMPRODUCT(--ISERROR({lrng('F')}))", 0),
        ("Errors in Payments flag and score columns", f"=SUMPRODUCT(--ISERROR(Payments!${PL['fd']}${pf}:${PL['pri']}${pl}))", 0),
        ("Errors in JE_Entries flag and score columns", f"=SUMPRODUCT(--ISERROR(JE_Entries!${JL['n']}${jf}:${JL['pri']}${jl}))", 0),
        ("High risk payments in list less High in Payments", f'=COUNTA(Exceptions!$A${r0 + 1}:$A${pend_row})-COUNTIF(Payments!${PL["pri"]}${pf}:${PL["pri"]}${pl},"High")', 0),
        ("High risk entries in list less High in JE_Entries", f'=COUNTA(Exceptions!$A${r1 + 1}:$A${jend_row})-COUNTIF(JE_Entries!${JL["pri"]}${jf}:${JL["pri"]}${jl},"High")', 0),
        ("Duplicate payment IDs", f"=SUMPRODUCT(--(COUNTIF(Payments!$A${pf}:$A${pl},Payments!$A${pf}:$A${pl})>1))", 0),
        ("Duplicate JE numbers in JE_Entries", f"=SUMPRODUCT(--(COUNTIF({jr('no')},{jr('no')})>1))", 0),
        ("Answer key rows with lookup errors", f"=SUMPRODUCT(--ISERROR({AK('F')}))", 0),
    ]
    for i, (lab, f, exp) in enumerate(ck):
        r = 5 + i
        ws.cell(row=r, column=1, value=lab)
        ws.cell(row=r, column=2, value=f)
        ws.cell(row=r, column=3, value=exp)
        ws.cell(row=r, column=4, value=f'=IF(ABS(B{r}-C{r})<0.005,"OK","FAIL")')
    r = 5 + len(ck) + 1
    ws.cell(row=r, column=1, value="OVERALL").font = BOLD
    ws.cell(row=r, column=4, value=f'=IF(COUNTIF(D5:D{4 + len(ck)},"FAIL")=0,"PASS","FAIL")').font = BOLD

    wb.save(path)
    return dict(pf=pf, pl=pl, jf=jf, jl=jl, PL=PL, JL=JL, ak_rows=len(ak_rows), n_high_p=len(hp), n_high_j=len(hj))


def summary_numbers(p, j, ben, vb, ak):
    out = {"payments": len(p), "entries": len(j), "benford": dict(n=ben["n"], mad=ben["mad"], chi=ben["chi"],
                                                                   conformity=conformity(ben["mad"]), counts=ben["counts"],
                                                                   z=ben["z"], sig=sum(1 for z in ben["z"] if z > PV["ben_z"]))}
    tests = {}
    pk = {"p_dup": "dup", "p_near": "near", "p_rnd": "rnd", "p_below": "below", "p_split": "split", "p_auth": "auth"}
    jk = {"j_" + k: k for k, _, _ in JE_W}
    for key, col in pk.items():
        tests[key] = int(p[col].sum())
    for key, col in jk.items():
        tests[key] = int(j[col if col != "susp" else "susp_f"].sum())
    tests["v_benf"] = sum(v["flag"] for v in vb.values())
    tests["v_emp"] = sum(v["emp"] for v in vb.values())
    out["flagged"] = tests
    out["high_pay"] = int((p.pri == "High").sum())
    out["high_je"] = int((j.pri == "High").sum())
    return out


def self_test(data, p, j, ben, vb):
    assert p.id.is_unique and j.no.is_unique
    assert abs(sum(sum(l[1] for l in e["lines"]) for e in data["entries"]) - sum(sum(l[2] for l in e["lines"]) for e in data["entries"])) > 0
    unb = j[(j.debit - j.credit).abs() > PV["tol"]]
    assert len(unb) == 3, len(unb)
    assert set(p.vendor) <= set(data["vendors"])
    assert all(CLASS[l[0]] for e in data["entries"] for l in e["lines"])
    return True


def main():
    out = sys.argv[1] if len(sys.argv) > 1 else "."
    os.makedirs(out, exist_ok=True)
    data = gen_data()
    p, j, ben, vb = pandas_tests(data)
    ak = build_answer_key(data, p, j, vb)
    self_test(data, p, j, ben, vb)
    path = os.path.join(out, "audit_analytics_workbook.xlsx")
    meta = build_workbook(data, p, j, ben, vb, ak, path)
    res = summary_numbers(p, j, ben, vb, ak)
    res["meta"] = {k: v for k, v in meta.items() if k not in ("PL", "JL")}
    with open(os.path.join(out, "results.json"), "w") as fh:
        json.dump(res, fh, indent=2, default=str)
    print("Workbook written to", path)
    print("Payments", len(p), "JE entries", len(j), "JE lines", sum(len(e["lines"]) for e in data["entries"]))
    print("Benford: n", ben["n"], "MAD", round(ben["mad"], 4), conformity(ben["mad"]), "chi", round(ben["chi"], 2))
    print("Flagged:", res["flagged"])
    print("High payments", res["high_pay"], "High entries", res["high_je"])


if __name__ == "__main__":
    main()
